#!/usr/bin/env python3
"""Build the independently updateable enemy configuration catalog."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


MAP_SHIP = re.compile(r"^\((\d+)\)(.+)$")
PARENS = str.maketrans("（）", "()")


def _rows(sheet):
    values = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(values)]
    for row_number, values_row in enumerate(values, start=2):
        yield row_number, {headers[index]: value for index, value in enumerate(values_row)}


def _text(value: Any) -> str | None:
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def _number(value: Any) -> int | float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else float(value)
    match = re.search(r"-?\d+(?:\.\d+)?", str(value).replace(",", ""))
    if match is None:
        return None
    number = float(match.group())
    return int(number) if number.is_integer() else number


def _stat(value: Any) -> dict[str, int | float] | None:
    if value is None or value == "":
        return None
    numbers = re.findall(r"-?\d+(?:\.\d+)?", str(value).replace(",", ""))
    if not numbers:
        return None
    parsed = [float(item) for item in numbers]
    parsed = [int(item) if item.is_integer() else item for item in parsed]
    result: dict[str, int | float] = {"base": parsed[0]}
    if len(parsed) > 1:
        result["equipped"] = parsed[1]
    return result


def _normalize(value: Any) -> str:
    return re.sub(r"[\s・･]", "", str(value or "").translate(PARENS)).lower()


def _key(row: dict[str, Any], row_number: int) -> str:
    identity = "|".join(
        str(row.get(name) or "")
        for name in ("No.", "敵艦名", "耐久", "火力・素（装備込）", "備考", "元テーブル")
    )
    digest = hashlib.sha1(identity.encode("utf-8")).hexdigest()[:10]
    return f"{int(row['No.'])}@{digest}"


SPECIAL_IDS = {
    "軽母ヌ級elite(艦載機白)": 1762,
    "軽母ヌ級elite(b)(艦載機白)": 1762,
    "軽母ヌ級elite(艦載機鳥白)": 1776,
    "軽母ヌ級elite(艦載機黒)": 1777,
    "軽母ヌ級flagship(艦載機白)": 1763,
    "軽母ヌ級flagship(b)(艦載機白)": 1763,
    "軽母ヌ級flagship(艦載機赤)": 1764,
    "軽母ヌ級flagship(c)(艦載機赤)": 1764,
    "軽母ヌ級改elite(艦載機鳥白)": 1765,
    "軽母ヌ級改elite(艦載機黒)": 1778,
    "軽母ヌ級改flagship(艦載機鳥赤)": 1766,
    "軽母ヌ級改flagship(艦載機赤)": 1779,
    "空母ヲ級flagship(艦載機白)": 1579,
    "空母ヲ級flagship(艦載機赤)": 1615,
    "空母ヲ級flagship(艦載機白赤)": 1614,
    "空母ヲ級改flagship(艦載機白)": 1616,
    "空母ヲ級改flagship(b)(艦載機白)": 1616,
    "空母ヲ級改flagship(艦載機赤)": 1618,
    "空母ヲ級改flagship(艦載機白赤)": 1617,
    "飛行場姫(陸爆弱)": 1650,
    "飛行場姫(陸爆中)": 1651,
    "飛行場姫(陸爆強)": 1652,
    "飛行場姫(鳥黒弱)": 2047,
    "飛行場姫(鳥黒強)": 2048,
    "飛行場姫(空襲)(a)": 1650,
    "飛行場姫(空襲)(f)": 2094,
    "飛行場姫(偵察)(a)": 1889,
    "港湾棲姫(最終形態)": 1613,
    "北方棲姫(前哨戦弱)": 1587,
    "北方棲姫(前哨戦強)": 1589,
    "北方棲姫(最終形態弱)": 1588,
    "北方棲姫(最終形態強)": 1590,
    "空母棲鬼(艦載機赤)": 1619,
    "空母棲鬼(艦載機白)": 1585,
    "空母棲姫(艦載機赤)": 1620,
    "空母棲姫(艦載機白)": 1586,
    "軽巡棲鬼(a)": 1601,
    "軽巡棲鬼(b)": 1602,
    "駆逐棲姫(a)": 1597,
    "駆逐棲姫(b)": 1598,
    "pt小鬼群(a)": 1637,
    "pt小鬼群(b)": 1638,
    "pt小鬼群(c)": 1639,
    "pt小鬼群(d)": 1640,
    "離島棲姫(a)": 1671,
    "離島棲姫(b)": 1672,
    "離島棲姫(陸爆弱)": 1668,
    "離島棲姫(陸爆強)": 1669,
    "砲台小鬼(a)": 1665,
    "砲台小鬼(b)": 1666,
    "砲台小鬼(c)": 1667,
    "潜水新棲姫(a)": 1736,
    "潜水新棲姫(b)": 1737,
    "潜水新棲姫(e)": 2049,
    "ヒ船団棲姫(a)": 2059,
    "ヒ船団棲姫(b)": 2060,
    "ヒ船団棲姫-壊(a)": 2061,
    "ヒ船団棲姫-壊(b)": 2062,
    "軽巡ヘ級改flagship(a)": 1904,
    "軽巡ヘ級改flagship(b)": 1905,
    "バタビア沖棲姫(a)": 1899,
    "バタビア沖棲姫(b)": 1900,
    "バタビア沖棲姫-壊(a)": 1902,
    "バタビア沖棲姫-壊(b)": 1903,
}
SPECIAL_IDS = {_normalize(key): value for key, value in SPECIAL_IDS.items()}


def _resolve_alias(raw_id: int, name: str, ships: list[dict[str, Any]]) -> tuple[str, str]:
    normalized = _normalize(name)
    wanted_id = SPECIAL_IDS.get(normalized)
    if wanted_id is not None:
        candidates = [ship for ship in ships if ship["id"] == wanted_id]
        if len(candidates) == 1:
            return candidates[0]["key"], "override"
        raise ValueError(
            f"Override for ({raw_id}){name} resolves to {len(candidates)} "
            f"configurations with id {wanted_id}"
        )

    exact = [
        ship
        for ship in ships
        if ship["id"] == raw_id and _normalize(ship["name"]) == normalized
    ]
    if len(exact) == 1:
        return exact[0]["key"], "exact-name"

    same_id = [ship for ship in ships if ship["id"] == raw_id]
    if raw_id == 1548 and normalized == _normalize("南方棲戦姫"):
        regular_map = [ship for ship in same_id if "5-3" in str(ship.get("note") or "")]
        if len(regular_map) == 1:
            return regular_map[0]["key"], "regular-map-override"
    if len(same_id) == 1:
        return same_id[0]["key"], "portrait-id"
    if len(exact) > 1:
        raise ValueError(f"Ambiguous exact-name match for ({raw_id}){name}")
    if same_id:
        raise ValueError(f"Ambiguous portrait-id match for ({raw_id}){name}")

    base = re.sub(r"\([^)]*\)", "", normalized)
    similar = [
        ship
        for ship in ships
        if ship["id"] == raw_id and _normalize(ship["name"]) == base
    ]
    if len(similar) == 1:
        return similar[0]["key"], "base-name"
    if similar:
        raise ValueError(f"Ambiguous base-name match for ({raw_id}){name}")
    raise ValueError(f"No enemy configuration matches ({raw_id}){name}")


def build_catalog(
    enemy_workbook: Path,
    sortie_workbook: Path,
    output: Path,
    *,
    data_version: str | None = None,
    revision: int = 1,
    published_at: str | None = None,
) -> dict[str, int]:
    workbook = load_workbook(enemy_workbook, read_only=True, data_only=True)
    required = {"enemy_ships", "enemy_equipment", "enemy_ship_equipment"}
    missing = required.difference(workbook.sheetnames)
    if missing:
        raise ValueError(f"Enemy workbook is missing sheets: {sorted(missing)}")

    equipment = {}
    for _, row in _rows(workbook["enemy_equipment"]):
        equipment_key = _text(row.get("装備キー"))
        if equipment_key is None:
            continue
        if equipment_key in equipment:
            workbook.close()
            raise ValueError(f"Duplicate enemy equipment key: {equipment_key}")
        equipment[equipment_key] = {
            "key": equipment_key,
            "name": _text(row.get("装備名")) or equipment_key,
            "type": _text(row.get("種別")),
            "firepower": _number(row.get("火力")),
            "torpedo": _number(row.get("雷装")),
            "bombing": _number(row.get("爆装")),
            "antiAir": _number(row.get("対空")),
            "antiSub": _number(row.get("対潜")),
            "search": _number(row.get("索敵")),
            "accuracy": _number(row.get("命中")),
            "evasion": _number(row.get("回避")),
            "armor": _number(row.get("装甲")),
            "range": _text(row.get("射程")),
            "note": _text(row.get("備考")),
        }
    equipment_by_name: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in equipment.values():
        equipment_by_name[_normalize(item["name"])].append(item)

    relation_rows = defaultdict(list)
    unmatched_equipment = 0
    for _, row in _rows(workbook["enemy_ship_equipment"]):
        relation_rows[(int(row["No."]), _normalize(row["敵艦名"]))].append(row)
        if _text(row.get("装備キー")) not in equipment:
            unmatched_equipment += 1

    ships = []
    for row_number, row in _rows(workbook["enemy_ships"]):
        if row.get("No.") is None:
            continue
        ship_id = int(row["No."])
        name = _text(row.get("敵艦名")) or f"Enemy {ship_id}"
        items = []
        related = relation_rows.get((ship_id, _normalize(name)), [])
        for slot in range(1, 6):
            equipment_name = _text(row.get(f"装備{slot}"))
            if equipment_name is None:
                continue
            candidates = equipment_by_name.get(_normalize(equipment_name), [])
            master = candidates[0] if len(candidates) == 1 else None
            if master is None:
                relation_matches = [
                    relation
                    for relation in related
                    if int(relation.get("装備スロット") or 0) == slot
                    and _normalize(relation.get("装備名（元表記）"))
                    == _normalize(equipment_name)
                ]
                relation_keys = {
                    _text(relation.get("装備キー")) for relation in relation_matches
                }
                relation_keys.discard(None)
                if len(relation_keys) == 1:
                    master = equipment.get(next(iter(relation_keys)))
            items.append({
                "slot": slot,
                "key": master.get("key") if master else None,
                "name": master.get("name") if master else equipment_name,
                "type": master.get("type") if master else _text(row.get(f"装備{slot}種別")),
                "matched": master is not None,
                "stats": master,
            })
        ships.append({
            "key": _key(row, row_number),
            "id": ship_id,
            "name": name,
            "shipType": _text(row.get("艦種")),
            "level": _number(row.get("Lv")),
            "hp": _number(row.get("耐久")),
            "firepower": _stat(row.get("火力・素（装備込）")),
            "torpedo": _stat(row.get("雷装・素（装備込）")),
            "antiAir": _stat(row.get("対空・素（装備込）")),
            "armor": _stat(row.get("装甲・素（装備込）")),
            "evasion": _number(row.get("回避")),
            "antiSub": _number(row.get("対潜")),
            "search": _number(row.get("索敵")),
            "luck": _number(row.get("運")),
            "aircraftCapacity": _number(row.get("搭載")),
            "speed": _text(row.get("速力")),
            "range": _text(row.get("射程")),
            "nightCutIn": _text(row.get("夜戦CI")),
            "note": _text(row.get("備考")),
            "detailsUrl": _text(row.get("詳細URL")),
            "equipment": items,
        })

    keys = [ship["key"] for ship in ships]
    if len(keys) != len(set(keys)):
        raise ValueError("Enemy configuration keys are not unique")

    sortie = load_workbook(sortie_workbook, read_only=True, data_only=True)
    if "敌舰配置" not in sortie.sheetnames:
        raise ValueError("Sortie workbook has no 敌舰配置 sheet")
    aliases: dict[str, str] = {}
    methods = defaultdict(int)
    for _, row in _rows(sortie["敌舰配置"]):
        for index in range(1, 13):
            raw = _text(row.get(f"敌舰{index}"))
            if raw is None:
                continue
            match = MAP_SHIP.match(raw)
            if match is None:
                raise ValueError(f"Invalid sortie enemy reference: {raw}")
            raw_id, name = int(match.group(1)), match.group(2).strip()
            alias = f"{raw_id}|{_normalize(name)}"
            resolved, method = _resolve_alias(raw_id, name, ships)
            previous = aliases.get(alias)
            if previous is not None and previous != resolved:
                raise ValueError(f"Conflicting sortie alias: {alias}")
            aliases[alias] = resolved
            methods[method] += 1

    timestamp = published_at or datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    catalog = {
        "schemaVersion": 1,
        "dataVersion": data_version or timestamp[:10],
        "revision": revision,
        "publishedAt": timestamp,
        "source": "kcwiki/wikiwiki enemy database",
        "ships": ships,
        "aliases": dict(sorted(aliases.items())),
        "quality": {
            "shipCount": len(ships),
            "uniqueEnemyIds": len({ship["id"] for ship in ships}),
            "equipmentCount": len(equipment),
            "mapAliasCount": len(aliases),
            "unmatchedEquipment": unmatched_equipment,
            "aliasMethods": dict(sorted(methods.items())),
        },
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(catalog, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    sortie.close()
    workbook.close()
    return {
        "ships": len(ships),
        "aliases": len(aliases),
        "unresolvedAliases": 0,
        "unmatchedEquipment": unmatched_equipment,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("enemy_workbook", type=Path)
    parser.add_argument("sortie_workbook", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--data-version")
    parser.add_argument("--revision", type=int, default=1)
    parser.add_argument("--published-at")
    args = parser.parse_args()
    report = build_catalog(
        args.enemy_workbook,
        args.sortie_workbook,
        args.output,
        data_version=args.data_version,
        revision=args.revision,
        published_at=args.published_at,
    )
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
