"""Build bundled sortie-map assets from the kcwiki workbook and image set."""

from __future__ import annotations

import argparse
import json
import re
import shutil
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from PIL import Image


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = PROJECT_ROOT / "data" / "sortie" / "source" / "kcwiki_出击_敌舰配置.xlsx"
DEFAULT_IMAGE_ROOT = PROJECT_ROOT / "data" / "sortie" / "source" / "images"
DEFAULT_OUTPUT_ROOT = PROJECT_ROOT / "assets"
DEFAULT_DATA_VERSION = "2026.09.20"
DEFAULT_REVISION = 2026092001
DEFAULT_PUBLISHED_AT = "2026-09-20T00:00:00Z"


SHIP_PATTERN = re.compile(r"^\((\d+)\)(.+)$")
ENEMY_CLASS_KANA = {
    "I": "イ",
    "RO": "ロ",
    "HA": "ハ",
    "NI": "ニ",
    "HO": "ホ",
    "HE": "ヘ",
    "TO": "ト",
    "CHI": "チ",
    "RI": "リ",
    "NU": "ヌ",
    "RU": "ル",
    "WO": "ヲ",
    "WA": "ワ",
    "KA": "カ",
    "YO": "ヨ",
    "TA": "タ",
    "RE": "レ",
    "SO": "ソ",
    "TSU": "ツ",
    "NE": "ネ",
}
ENEMY_CLASS_PATTERN = re.compile(
    rf"(?<![A-Z])({'|'.join(sorted(ENEMY_CLASS_KANA, key=len, reverse=True))})(?=级)"
)


def _rows(sheet: Any) -> Iterable[dict[str, Any]]:
    values = sheet.iter_rows(values_only=True)
    headers = [str(value) if value is not None else "" for value in next(values)]
    for row in values:
        if not any(value is not None for value in row):
            continue
        yield dict(zip(headers, row))


def _map_sort_key(map_id: str) -> tuple[int, int]:
    area, number = map_id.split("-", maxsplit=1)
    return int(area), int(number)


def _difficulty(value: Any) -> int:
    return str(value or "").count("☆")


def _japanese_map_name(value: Any) -> str:
    return str(value or "").translate(str.maketrans({"冲": "沖", "帶": "帯"}))


def _enemy_name_zh(value: Any) -> str:
    name = str(value or "")
    name = ENEMY_CLASS_PATTERN.sub(
        lambda match: ENEMY_CLASS_KANA[match.group(1)],
        name,
    )
    name = re.sub(r"(elite|flagship)\s+e(?=\s)", r"\1", name)
    name = re.sub(r"级(?=(?:elite|flagship))", "级 ", name)
    return re.sub(r"(elite|flagship)(?=[\u4e00-\u9fff])", r"\1 ", name)


def _truthy(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() not in {"", "0", "false", "no", "否", "无"}


def _optional_int(value: Any) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, bool):
        raise ValueError(f"Expected an integer, got {value!r}")
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = str(value).strip()
    if re.fullmatch(r"[+-]?\d+", text):
        return int(text)
    raise ValueError(f"Expected an integer, got {value!r}")


def _node_kind(
    label: str,
    raw: Any,
    has_formations: bool,
    battle_type: Any = None,
) -> str:
    raw_text = str(raw or "").strip().lower()
    if raw_text:
        return raw_text
    if "BOSS" in label.upper():
        return "boss"
    if label == "ボス":
        return "boss"
    if label == "資源":
        return "resource"
    if label == "空襲":
        return "airstrike"
    if label == "能動分岐":
        return "activebranch"
    if label == "戦闘回避":
        return "imaginary"
    if label == "旋流":
        return "maelstorm1"
    if str(battle_type or "").strip() == "夜戦":
        return "night"
    if "资源" in label:
        return "resource"
    if "航空" in label:
        return "air"
    if "夜战" in label:
        return "night"
    if has_formations or "战斗" in label:
        return "battle"
    return "other"


def _ship_entry(value: Any, ship_names: dict[str, dict[str, str]]) -> dict[str, Any] | None:
    text = str(value or "").strip()
    if not text:
        return None
    match = SHIP_PATTERN.match(text)
    if match is None:
        return {"id": None, "nameJa": text, "nameZh": text}
    ship_id, fallback_name = match.groups()
    names = ship_names.get(ship_id, {})
    return {
        "id": int(ship_id),
        # The same portrait id is shared by multiple configurations. The
        # per-formation label is therefore more precise than the id lookup.
        "nameJa": fallback_name or names.get("nameJa", ""),
        "nameZh": fallback_name or names.get("nameZh", ""),
    }


def _image_for(directory: Path, map_id: str) -> Path:
    matches = sorted(directory.glob(f"{map_id}*.png"))
    if len(matches) != 1:
        raise ValueError(
            f"Expected exactly one PNG for {map_id} in {directory}, found {len(matches)}"
        )
    return matches[0]


def _copy_trimmed_png(source: Path, target: Path) -> float:
    """Copy a PNG after removing its fully transparent outer rows and columns."""
    with Image.open(source) as image:
        if (
            image.format != "PNG"
            or image.width <= 0
            or image.height <= 0
            or image.width > 4096
            or image.height > 4096
            or image.width * image.height > 8_000_000
        ):
            raise ValueError(f"Invalid or oversized PNG: {source}")
        rgba = image.convert("RGBA")
        alpha_bounds = rgba.getchannel("A").getbbox()
        trimmed = rgba.crop(alpha_bounds) if alpha_bounds is not None else rgba
        trimmed.save(target)
        return trimmed.width / trimmed.height


def _copy_validated_png(source: Path, target: Path) -> None:
    """Validate the complete PNG payload and copy it to the asset snapshot."""
    with Image.open(source) as image:
        width, height = image.size
        if (
            image.format != "PNG"
            or width <= 0
            or height <= 0
            or width > 4096
            or height > 4096
            or width * height > 8_000_000
        ):
            raise ValueError(f"Invalid or oversized PNG: {source}")
        image.load()
    shutil.copy2(source, target)


def build_assets(
    workbook_path: Path,
    image_root: Path,
    output_root: Path,
    *,
    data_version: str = DEFAULT_DATA_VERSION,
    revision: int = DEFAULT_REVISION,
    published_at: str = DEFAULT_PUBLISHED_AT,
) -> dict[str, Any]:
    existing_catalog_path = output_root / "data" / "sortie_map_catalog.json"
    existing_maps: dict[str, dict[str, Any]] = {}
    if existing_catalog_path.is_file():
        existing_catalog = json.loads(existing_catalog_path.read_text(encoding="utf-8"))
        for item in existing_catalog.get("maps", []):
            map_id = str(item.get("id") or "")
            if not map_id:
                continue
            existing_maps[map_id] = item

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    required_sheets = {"敌舰配置", "节点索引", "海域索引", "舰船ID对照"}
    missing = required_sheets.difference(workbook.sheetnames)
    if missing:
        raise ValueError(f"Workbook is missing sheets: {sorted(missing)}")

    ship_names: dict[str, dict[str, str]] = {}
    for row in _rows(workbook["舰船ID对照"]):
        ship_id = str(row["舰船ID"])
        ship_names[ship_id] = {
            "nameJa": str(row.get("日文名") or ""),
            "nameZh": _enemy_name_zh(row.get("中文名") or row.get("日文名")),
        }

    formations_by_node: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in _rows(workbook["敌舰配置"]):
        map_id = str(row["海域编号"])
        point = str(row["点"])
        ships = [
            entry
            for index in range(1, 13)
            if (entry := _ship_entry(row.get(f"敌舰{index}"), ship_names)) is not None
        ]
        first_group_size = len(ships)
        grouping = str(row.get("编队分组") or "")
        if "+" in grouping:
            first_text, _, _ = grouping.partition("+")
            try:
                first_group_size = max(0, min(len(ships), int(first_text)))
            except ValueError:
                first_group_size = len(ships)
        groups = [ships[:first_group_size]]
        if first_group_size < len(ships):
            groups.append(ships[first_group_size:])
        formations_by_node[(map_id, point)].append(
            {
                "variant": int(row.get("变体") or 1),
                "final": _truthy(row.get("最终形态") or row.get("斩杀配置")),
                "formation": row.get("阵型"),
                "experience": _optional_int(row.get("经验值")),
                "airPower": _optional_int(row.get("制空值")),
                "airSuperiority": _optional_int(row.get("空优值")),
                "airSupremacy": _optional_int(row.get("空确值")),
                "fleetGroups": groups,
                "note": row.get("备注"),
            }
        )

    nodes_by_map: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in _rows(workbook["节点索引"]):
        map_id = str(row["海域编号"])
        point = str(row["点"])
        formations = formations_by_node.get((map_id, point), [])
        label = str(row.get("点类型") or "")
        battle_type = str(row.get("战斗类型") or "")
        detected_kind = _node_kind(
            label,
            row.get("点类型原始"),
            bool(formations),
            battle_type,
        )
        nodes_by_map[map_id].append(
            {
                "point": point,
                "kind": detected_kind,
                "typeLabel": label,
                "battleTypeLabel": battle_type,
                "nameJa": row.get("点原名"),
                "reward": row.get("奖励原始") or row.get("奖励资源"),
                "formations": formations,
            }
        )

    cover_output = output_root / "images" / "sortie_maps" / "covers"
    map_output = output_root / "images" / "sortie_maps" / "maps"
    data_output = output_root / "data"
    cover_output.mkdir(parents=True, exist_ok=True)
    map_output.mkdir(parents=True, exist_ok=True)
    data_output.mkdir(parents=True, exist_ok=True)

    maps: list[dict[str, Any]] = []
    map_rows = list(_rows(workbook["海域索引"]))
    for row in sorted(
        map_rows,
        key=lambda item: _map_sort_key(str(item.get("海域编号") or item.get("map"))),
    ):
        map_id = str(row.get("海域编号") or row.get("map"))
        existing_map = existing_maps.get(map_id, {})
        cover_source = _image_for(image_root / "封面", map_id)
        map_source = _image_for(image_root / "路线图", map_id)
        cover_target = cover_output / f"{map_id}.png"
        map_target = map_output / f"{map_id}.png"
        _copy_validated_png(cover_source, cover_target)
        map_aspect_ratio = _copy_trimmed_png(map_source, map_target)
        maps.append(
            {
                "id": map_id,
                "nameJa": _japanese_map_name(row.get("海域原名") or row.get("name_ja")),
                "difficulty": _difficulty(row.get("difficulty"))
                if row.get("difficulty") is not None
                else int(existing_map.get("difficulty") or 0),
                "coverAsset": f"assets/images/sortie_maps/covers/{map_id}.png",
                "mapAsset": f"assets/images/sortie_maps/maps/{map_id}.png",
                "mapAspectRatio": map_aspect_ratio,
                "source": row.get("来源"),
                "nodes": nodes_by_map.get(map_id, []),
            }
        )

    catalog = {
        "version": 1,
        "schemaVersion": 1,
        "dataVersion": data_version,
        "revision": revision,
        "publishedAt": published_at,
        "source": "https://zh.kcwiki.cn/wiki/出击",
        "maps": maps,
    }
    workbook.close()
    catalog_path = data_output / "sortie_map_catalog.json"
    catalog_path.write_text(
        json.dumps(catalog, ensure_ascii=False, separators=(",", ":")) + "\n",
        encoding="utf-8",
    )
    return catalog


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path, nargs="?", default=DEFAULT_WORKBOOK)
    parser.add_argument("image_root", type=Path, nargs="?", default=DEFAULT_IMAGE_ROOT)
    parser.add_argument("output_root", type=Path, nargs="?", default=DEFAULT_OUTPUT_ROOT)
    parser.add_argument("--data-version", default=DEFAULT_DATA_VERSION)
    parser.add_argument("--revision", type=int, default=DEFAULT_REVISION)
    parser.add_argument("--published-at", default=DEFAULT_PUBLISHED_AT)
    args = parser.parse_args()
    catalog = build_assets(
        args.workbook,
        args.image_root,
        args.output_root,
        data_version=args.data_version,
        revision=args.revision,
        published_at=args.published_at,
    )
    print(
        json.dumps(
            {
                "maps": len(catalog["maps"]),
                "nodes": sum(len(item["nodes"]) for item in catalog["maps"]),
                "formations": sum(
                    len(node["formations"])
                    for item in catalog["maps"]
                    for node in item["nodes"]
                ),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
