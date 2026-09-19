import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from build_enemy_catalog import SPECIAL_IDS, _normalize, _resolve_alias, build_catalog


class BuildEnemyCatalogTest(unittest.TestCase):
    def test_ambiguous_configuration_requires_an_explicit_override(self):
        ships = [
            {"id": 100, "name": "重複名", "key": "100@a"},
            {"id": 100, "name": "重複名", "key": "100@b"},
        ]

        with self.assertRaisesRegex(ValueError, "Ambiguous exact-name"):
            _resolve_alias(100, "重複名", ships)

    def test_non_override_never_maps_to_a_different_portrait_id(self):
        ships = [{"id": 101, "name": "同名敵艦", "key": "101@only"}]

        with self.assertRaisesRegex(ValueError, "No enemy configuration"):
            _resolve_alias(100, "同名敵艦", ships)

    def test_duplicate_equipment_key_is_rejected(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            enemy = root / "enemy.xlsx"
            sortie = root / "sortie.xlsx"
            self._write_enemy(enemy)
            self._write_sortie(sortie)
            workbook = load_workbook(enemy)
            workbook["enemy_equipment"].append(
                ["5inch", "重複装備", "小口径主砲"]
            )
            workbook.save(enemy)
            workbook.close()

            with self.assertRaisesRegex(ValueError, "Duplicate enemy equipment key"):
                build_catalog(enemy, sortie, root / "catalog.json")

    def test_known_portrait_variants_use_their_real_configuration_ids(self):
        expected = {
            "空母棲姫(艦載機白)": 1586,
            "潜水新棲姫(B)": 1737,
            "離島棲姫(陸爆強)": 1669,
            "ヒ船団棲姫-壊(B)": 2062,
        }
        self.assertEqual(
            {name: SPECIAL_IDS[_normalize(name)] for name in expected}, expected
        )

    def test_builds_unique_configurations_equipment_and_map_aliases(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            enemy = root / "enemy.xlsx"
            sortie = root / "sortie.xlsx"
            output = root / "enemy_catalog.json"
            self._write_enemy(enemy)
            self._write_sortie(sortie)

            report = build_catalog(enemy, sortie, output, revision=7)
            catalog = json.loads(output.read_text(encoding="utf-8"))

            self.assertEqual(catalog["revision"], 7)
            self.assertEqual(len(catalog["ships"]), 2)
            self.assertEqual(catalog["ships"][0]["equipment"][0]["name"], "5inch単装砲")
            alias = catalog["aliases"]["1523|軽母ヌ級elite(艦載機黒)"]
            resolved = next(ship for ship in catalog["ships"] if ship["key"] == alias)
            self.assertEqual(resolved["id"], 1777)
            self.assertEqual(report["unresolvedAliases"], 0)

    @staticmethod
    def _write_enemy(path: Path):
        workbook = Workbook()
        ships = workbook.active
        ships.title = "enemy_ships"
        ships.append([
            "No.", "敵艦名", "艦種", "Lv", "耐久", "火力・素（装備込）",
            "雷装・素（装備込）", "対空・素（装備込）", "装甲・素（装備込）",
            "回避", "対潜", "索敵", "運", "搭載", "速力", "射程",
            "装備1", "装備1種別", "装備2", "装備2種別", "装備3", "装備3種別",
            "装備4", "装備4種別", "装備5", "装備5種別", "夜戦CI", "備考",
            "詳細URL", "元テーブル",
        ])
        ships.append([1523, "軽母ヌ級 elite", "軽空母", 1, 48, "10（18）", 0, 20, 25,
                      10, 0, 10, 1, 18, "低速", "中", "5inch単装砲", "小口径主砲"])
        ships.append([1777, "軽母ヌ級 elite", "軽空母", 1, 48, "10（18）", 0, 20, 25,
                      10, 0, 10, 1, 18, "低速", "中", "5inch単装砲", "小口径主砲",
                      None, None, None, None, None, None, None, "艦載機黒"])
        equipment = workbook.create_sheet("enemy_equipment")
        equipment.append(["装備キー", "装備名", "種別", "火力", "雷装", "爆装", "対空",
                          "対潜", "索敵", "命中", "回避", "装甲", "射程", "備考"])
        equipment.append(["5inch", "5inch単装砲", "小口径主砲", 1, 0, 0, 0, 0, 0, 0, 0, 0, "短", None])
        relations = workbook.create_sheet("enemy_ship_equipment")
        relations.append(["No.", "敵艦名", "装備スロット", "装備名（元表記）", "装備名",
                          "装備キー", "種別", "マッチ状態", "元テーブル"])
        relations.append([1523, "軽母ヌ級 elite", 1, "5inch単装砲", "5inch単装砲", "5inch", "小口径主砲", "matched", "x"])
        relations.append([1777, "軽母ヌ級 elite", 1, "5inch単装砲", "5inch単装砲", "5inch", "小口径主砲", "matched", "x"])
        workbook.save(path)

    @staticmethod
    def _write_sortie(path: Path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "敌舰配置"
        sheet.append(["海域编号", "点", "敌舰1"])
        sheet.append(["1-1", "A", "(1523)軽母ヌ級elite(艦載機黒)"])
        workbook.save(path)


if __name__ == "__main__":
    unittest.main()
