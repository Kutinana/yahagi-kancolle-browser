import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from PIL import Image

from tool.build_sortie_map_assets import (
    _copy_trimmed_png,
    _copy_validated_png,
    _enemy_name_zh,
    build_assets,
)


class BuildSortieMapAssetsTest(unittest.TestCase):
    def test_rejects_a_truncated_cover_png(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "broken.png"
            target = root / "output.png"
            source.write_bytes(b"\x89PNG\r\n\x1a\ntruncated")

            with self.assertRaises(Exception):
                _copy_validated_png(source, target)
            self.assertFalse(target.exists())

    def test_normalizes_enemy_class_kana_and_known_spacing_artifacts(self):
        self.assertEqual(_enemy_name_zh("空母WO级"), "空母ヲ级")
        self.assertEqual(_enemy_name_zh("未知HI级"), "未知HI级")
        self.assertEqual(
            _enemy_name_zh("轻母NU级改 elite e 舰载机(鸟白)"),
            "轻母ヌ级改 elite 舰载机(鸟白)",
        )
        self.assertEqual(
            _enemy_name_zh("空母WO级改 flagship舰载机（白赤）"),
            "空母ヲ级改 flagship 舰载机（白赤）",
        )
        self.assertEqual(
            _enemy_name_zh("轻巡洋舰HE级flagship"),
            "轻巡洋舰ヘ级 flagship",
        )

    def test_fully_transparent_route_map_keeps_its_original_extent(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "transparent.png"
            target = root / "output.png"
            Image.new("RGBA", (9, 5), (0, 0, 0, 0)).save(source)

            self.assertEqual(_copy_trimmed_png(source, target), 9 / 5)
            with Image.open(target) as image:
                self.assertEqual(image.mode, "RGBA")
                self.assertEqual(image.size, (9, 5))

    def test_rgb_route_map_is_converted_without_losing_its_extent(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "rgb.png"
            target = root / "output.png"
            Image.new("RGB", (7, 3), (20, 40, 60)).save(source)

            self.assertEqual(_copy_trimmed_png(source, target), 7 / 3)
            with Image.open(target) as image:
                self.assertEqual(image.mode, "RGBA")
                self.assertEqual(image.size, (7, 3))

    def test_builds_japanese_map_catalog_and_stable_image_names(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            workbook_path = root / "source.xlsx"
            image_root = root / "images"
            cover_dir = image_root / "封面"
            map_dir = image_root / "路线图"
            cover_dir.mkdir(parents=True)
            map_dir.mkdir(parents=True)
            Image.new("RGBA", (4, 2), (20, 40, 60, 255)).save(
                cover_dir / "1-1中文名.png"
            )
            route_image = Image.new("RGBA", (10, 6), (0, 0, 0, 0))
            route_image.paste((20, 40, 60, 255), (1, 1, 9, 5))
            route_image.save(map_dir / "1-1中文名+全图.png")

            workbook = Workbook()
            enemy_sheet = workbook.active
            enemy_sheet.title = "敌舰配置"
            enemy_sheet.append(
                [
                    "海域编号",
                    "海域原名",
                    "海域译名",
                    "点",
                    "点类型",
                    "战斗类型",
                    "点原名",
                    "点译名",
                    "变体",
                    "最终形态",
                    "阵型",
                    "阵型原始",
                    "经验值",
                    "制空值",
                    "空优值",
                    "空确值",
                    "敌舰1",
                    "敌舰2",
                    "敌舰3",
                    "敌舰4",
                    "敌舰5",
                    "敌舰6",
                    "敌舰7",
                    "敌舰8",
                    "敌舰9",
                    "敌舰10",
                    "敌舰11",
                    "敌舰12",
                    "舰船数量",
                    "编队分组",
                    "奖励资源",
                    "奖励原始",
                    "备注",
                ]
            )
            enemy_sheet.append(
                [
                    "1-1",
                    "製油所地帶沿岸",
                    "中文名",
                    "C",
                    "ボス",
                    "通常戦闘",
                    "敵主力艦隊",
                    None,
                    1,
                    "是",
                    "単縦陣",
                    "単縦陣",
                    "50",
                    "27",
                    "42",
                    "82",
                    "(1512)空母ヲ級",
                    "(1501)駆逐イ級",
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    2,
                    None,
                    None,
                    None,
                    None,
                ]
            )

            node_sheet = workbook.create_sheet("节点索引")
            node_sheet.append(
                [
                    "海域编号",
                    "海域原名",
                    "海域译名",
                    "点",
                    "点类型",
                    "战斗类型",
                    "点原名",
                    "点译名",
                    "变体数",
                    "战斗变体数",
                    "敌舰数量合计",
                    "奖励资源",
                    "奖励原始",
                ]
            )
            node_sheet.append(
                [
                    "1-1",
                    "製油所地帶沿岸",
                    "中文名",
                    "C",
                    "ボス",
                    "通常戦闘",
                    "敵主力艦隊",
                    None,
                    1,
                    1,
                    2,
                    None,
                    None,
                ]
            )

            map_sheet = workbook.create_sheet("海域索引")
            map_sheet.append(
                [
                    "海域编号",
                    "海域原名",
                    "来源",
                ]
            )
            map_sheet.append(
                [
                    "1-1",
                    "製油所地帶沿岸",
                    "https://example.test/1-1",
                ]
            )

            ship_sheet = workbook.create_sheet("舰船ID对照")
            ship_sheet.append(["舰船ID", "日文名", "出现次数"])
            ship_sheet.append(["1501", "駆逐イ級", 1])
            ship_sheet.append(["1505", "軽巡ホ級", 1])
            ship_sheet.append(["1512", "空母ヲ級", 1])
            workbook.save(workbook_path)

            output_root = root / "assets"
            catalog_path = output_root / "data" / "sortie_map_catalog.json"
            catalog_path.parent.mkdir(parents=True)
            catalog_path.write_text(
                json.dumps(
                    {
                        "maps": [
                            {
                                "id": "1-1",
                                "difficulty": 2,
                                "nodes": [{"point": "C", "kind": "未知"}],
                            }
                        ]
                    }
                ),
                encoding="utf-8",
            )
            build_assets(
                workbook_path,
                image_root,
                output_root,
                data_version="2026.09.20",
                revision=2026092001,
                published_at="2026-09-20T00:00:00Z",
            )

            catalog = json.loads(catalog_path.read_text(encoding="utf-8"))
            self.assertEqual(catalog["version"], 1)
            self.assertEqual(catalog["schemaVersion"], 1)
            self.assertEqual(catalog["dataVersion"], "2026.09.20")
            self.assertEqual(catalog["revision"], 2026092001)
            self.assertEqual(catalog["publishedAt"], "2026-09-20T00:00:00Z")
            self.assertEqual(len(catalog["maps"]), 1)
            map_info = catalog["maps"][0]
            self.assertEqual(map_info["id"], "1-1")
            self.assertEqual(map_info["nameJa"], "製油所地帯沿岸")
            self.assertEqual(map_info["difficulty"], 2)
            self.assertEqual(map_info["coverAsset"], "assets/images/sortie_maps/covers/1-1.png")
            self.assertEqual(map_info["mapAsset"], "assets/images/sortie_maps/maps/1-1.png")
            self.assertEqual(map_info["mapAspectRatio"], 2.0)
            formation = map_info["nodes"][0]["formations"][0]
            self.assertEqual(map_info["nodes"][0]["kind"], "boss")
            self.assertEqual(map_info["nodes"][0]["typeLabel"], "ボス")
            self.assertEqual(map_info["nodes"][0]["battleTypeLabel"], "通常戦闘")
            self.assertTrue(formation["final"])
            self.assertEqual(formation["experience"], 50)
            self.assertEqual(formation["airPower"], 27)
            self.assertEqual(formation["airSuperiority"], 42)
            self.assertEqual(formation["airSupremacy"], 82)
            self.assertEqual(formation["fleetGroups"][0][0]["nameJa"], "空母ヲ級")
            self.assertTrue((output_root / "images/sortie_maps/covers/1-1.png").is_file())
            self.assertTrue((output_root / "images/sortie_maps/maps/1-1.png").is_file())
            with Image.open(output_root / "images/sortie_maps/maps/1-1.png") as image:
                self.assertEqual(image.size, (8, 4))


if __name__ == "__main__":
    unittest.main()
